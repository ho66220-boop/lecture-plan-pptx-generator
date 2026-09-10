# -*- coding: utf-8 -*-
"""2026-09 전체 구현 리뷰에서 확인된 문제들의 회귀 테스트.

각 테스트는 리뷰 당시 재현된 '조용한 오생성' 사례를 그대로 고정한다 — 값이 틀리는데
리포트가 비어 있던 경로에 값 수정 또는 리포트를 붙였다.
"""
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Cm

from src import fee as fee_mod
from src.collect_excel import collect_lectures
from src.date_utils import extract_dates, extract_start_time, extract_start_time_display, parse_date_detail
from src.export_outputs import export_validation_report
from src.generate_pptx import (
    generate_pptx_from_template,
    progress_split,
    template_progress_capacity,
    validate_progress_overflow,
)
from src.normalize import is_holiday_row, load_academic_calendar, normalize_lecture, normalize_lectures
from src.pipeline import run_pipeline

PHOTO_DIR = Path("teacher_photos")


def make_raw(dates, **fields_override):
    fields = {
        "강사명": "테스트강사", "강좌명": "A고 국어 내신 대비반", "메인 제목": "테스트", "과목": "국어",
        "구분": "정규반", "강의형태": "현장강의", "수업 요일 / 시간": "화 19:00",
    }
    fields.update(fields_override)
    return {
        "source_file": "s.xlsx", "source_sheet": "강좌1", "fields": fields,
        "progress": [{"날짜": d, "수업 주제": f"{i + 1}강"} for i, d in enumerate(dates)],
    }


def codes(reports):
    return [r["issue_code"] for r in reports]


def write_sheet(path, raw, progress_rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "강좌1"
    for key, value in raw["fields"].items():
        ws.append([key, value])
    ws.append(["회차", "날짜", "수업 주제", "상세 내용", "비고"])
    for row in progress_rows if progress_rows is not None else raw["progress"]:
        ws.append([row.get(k, "") for k in ("회차", "날짜", "수업 주제", "상세 내용", "비고")])
    wb.save(path)
    return path


# ── High 1: 연도가 명시된 날짜에 연말 보정을 중복 적용하지 않는다 ──

def test_explicit_year_not_bumped_after_year_wrap():
    lec, reports = normalize_lecture(make_raw(["2026-12-28", "2027-01-04", "2027-01-11"]), 1, 2026)
    assert [r["parsed_date"] for r in lec["progress"]] == ["2026-12-28", "2027-01-04", "2027-01-11"]
    assert lec["computed_total_sessions"] == 3
    assert "OUT_OF_ORDER_DATE" not in codes(reports)


def test_mixed_explicit_and_implicit_years_stay_consistent():
    lec, _ = normalize_lecture(make_raw(["12/28", "1/4", "2027-01-11"]), 1, 2026)
    assert [r["parsed_date"] for r in lec["progress"]] == ["2026-12-28", "2027-01-04", "2027-01-11"]


def test_december_preview_keeps_explicit_january():
    lec, _ = normalize_lecture(make_raw(["2026-12-28", "2027-01-04"]), 1, 2026, target_month=12)
    assert [r["parsed_date"] for r in lec["progress"]] == ["2026-12-28", "2027-01-04"]
    assert lec["progress"][1]["is_next_month"] is True
    assert lec["computed_total_sessions"] == 1


def test_parse_date_detail_reports_explicit_year():
    assert parse_date_detail("2027-01-04")[2] is True
    assert parse_date_detail("1/4", 2026)[2] is False
    assert parse_date_detail(date(2027, 1, 4))[2] is True


# ── High 2: 공백이 있는 FEE_OVERRIDES 설정 키도 매칭되고, 미사용 키는 리포트 ──

def test_override_key_with_spaces_matches(monkeypatch):
    monkeypatch.setitem(fee_mod.FEE_OVERRIDES, "테스트강사|A고 국어 내신 대비반", 123456)
    lec, _ = normalize_lecture(make_raw(["7/7", "7/14", "7/21", "7/28"]), 1, 2026)
    assert lec["computed_fee"] == 123456
    assert lec["billing"] == "override"


def test_unused_override_key_reported(monkeypatch):
    monkeypatch.setitem(fee_mod.FEE_OVERRIDES, "없는강사|없는강좌", 1)
    _, reports = normalize_lectures([make_raw(["7/7"])], 2026)
    unused = [r for r in reports if r["issue_code"] == "FEE_OVERRIDE_UNUSED"]
    assert len(unused) == 1 and unused[0]["severity"] == "정보"
    assert "없는강사|없는강좌" in unused[0]["message"]


def test_matched_override_key_not_reported(monkeypatch):
    monkeypatch.setitem(fee_mod.FEE_OVERRIDES, "테스트강사|A고 국어 내신 대비반", 1)
    _, reports = normalize_lectures([make_raw(["7/7"])], 2026)
    assert "FEE_OVERRIDE_UNUSED" not in codes(reports)


# ── High 3: 휴강 판정 — 부정 표현 무시, 휴강일 칸과 진도표 대조 ──

@pytest.mark.parametrize("text,expected", [
    ("휴강", True), ("7/14 휴강", True), ("내신기간 휴강", True),
    ("휴강 없음", False), ("휴강X", False), ("휴강 아님", False), ("휴강 없이 진행", False),
])
def test_holiday_negation(text, expected):
    assert is_holiday_row({"비고": text}) is expected


def test_holiday_negation_row_counts_as_class():
    raw = make_raw(["7/7", "7/14"])
    raw["progress"][1]["비고"] = "휴강 없음"
    lec, _ = normalize_lecture(raw, 1, 2026)
    assert lec["computed_total_sessions"] == 2


def test_holiday_field_date_counted_as_class_is_reported():
    raw = make_raw(["7/7", "7/14"], **{"휴강일 / 사유 / 수업 불가 일정": "7/14 휴강"})
    lec, reports = normalize_lecture(raw, 1, 2026)
    assert lec["computed_total_sessions"] == 2          # 값은 고치지 않는다(정책)
    conflict = [r for r in reports if r["issue_code"] == "HOLIDAY_FIELD_CONFLICT"]
    assert len(conflict) == 1 and "2행" in conflict[0]["message"]
    assert "HOLIDAY_FIELD_CONFLICT" in lec["flags"]


def test_progress_holiday_missing_from_field_is_reported():
    raw = make_raw(["7/7", "7/14"], **{"휴강일 / 사유 / 수업 불가 일정": "7/21 휴강"})
    raw["progress"][1]["비고"] = "휴강"
    lec, reports = normalize_lecture(raw, 1, 2026)
    assert lec["computed_total_sessions"] == 1
    assert codes(reports).count("HOLIDAY_FIELD_CONFLICT") == 1


def test_holiday_field_without_dates_is_not_cross_checked():
    raw = make_raw(["7/7", "7/14"], **{"휴강일 / 사유 / 수업 불가 일정": "내신기간 휴강"})
    raw["progress"][1]["비고"] = "휴강"
    _, reports = normalize_lecture(raw, 1, 2026)
    assert "HOLIDAY_FIELD_CONFLICT" not in codes(reports)


def test_extract_dates_handles_iso_and_slash():
    assert extract_dates("7/14, 7/21 내신 휴강 2026-08-04(화)", 2026) == [
        date(2026, 8, 4), date(2026, 7, 14), date(2026, 7, 21),
    ]


# ── High 4: 오전/오후 반영 + 시·분 범위 검증 ──

@pytest.mark.parametrize("text,expected", [
    ("매주 월요일 18:30~22:00", "오후 6시 30분"),
    ("화 오후 2:00~5:00", "오후 2시"),
    ("화 오후 7시~10시", "오후 7시"),
    ("오전 12:30", "오전 12시 30분"),
    ("2:00 PM", "오후 2시"),
    ("오후 14:00", "오후 2시"),
    ("목 9시 30분", "오전 9시 30분"),
    ("화 25:99", None),
    ("화 19:60", None),
])
def test_extract_start_time_display(text, expected):
    assert extract_start_time_display(text) == expected


def test_ambiguous_small_hour_flagged():
    assert extract_start_time("화 2:00~5:00")["ambiguous"] is True
    assert extract_start_time("화 오후 2:00")["ambiguous"] is False
    assert extract_start_time("화 19:00")["ambiguous"] is False
    _, reports = normalize_lecture(make_raw(["7/7"], **{"수업 요일 / 시간": "화 2:00~5:00"}), 1, 2026)
    assert "OPENING_TIME_AMBIGUOUS" in codes(reports)


def test_out_of_range_time_reported_not_displayed():
    lec, reports = normalize_lecture(make_raw(["7/7"], **{"수업 요일 / 시간": "화 25:99"}), 1, 2026)
    assert lec["opening_date_display"] == "07. 07(화)"
    assert "OPENING_TIME_PARSE_FAILED" in codes(reports)


# ── High 5: 정규화 산출 저장이 실패해도 이번 실행의 리포트는 저장된다 ──

def test_report_written_even_if_normalized_export_fails(tmp_path):
    src = write_sheet(tmp_path / "in.xlsx", make_raw(["7/7"]))
    out = tmp_path / "out"
    out.mkdir()
    (out / "validation_report.xlsx").write_bytes(b"previous-run")
    (out / "normalized_data.json").mkdir()          # JSON 저장을 강제로 실패시킨다
    with pytest.raises(Exception):
        run_pipeline(src, out, make_pptx=False)
    assert (out / "validation_report.xlsx").read_bytes() != b"previous-run"
    load_workbook(out / "validation_report.xlsx")     # 유효한 엑셀


# ── Medium: 대상 월 검증, 시즌=정규반 필터 일관성 ──

def test_invalid_target_month_rejected(tmp_path):
    src = write_sheet(tmp_path / "in.xlsx", make_raw(["7/7"]))
    with pytest.raises(ValueError):
        run_pipeline(src, tmp_path / "out", make_pptx=False, target_month=13)
    with pytest.raises(ValueError):
        normalize_lectures([make_raw(["7/7"])], 2026, target_month=0)


def test_missing_input_path_fails_loudly(tmp_path):
    with pytest.raises(FileNotFoundError):
        run_pipeline(tmp_path / "없는파일.xlsx", tmp_path / "out", make_pptx=False)


def test_season_regular_is_filtered_by_month():
    lec, _ = normalize_lecture(make_raw(["7/7", "8/4"], 구분="단과", 시즌="정규반"), 1, 2026, target_month=7)
    assert lec["billing"] == "monthly"
    assert lec["computed_total_sessions"] == 1
    assert lec["progress"][1]["is_next_month"] is True


# ── Medium: 학사일정 검사가 돌았는지 리포트로 알 수 있다 ──

def test_calendar_missing_file_reported(tmp_path):
    events, reports = load_academic_calendar(tmp_path / "none.csv", base_year=2026)
    assert events == {} and codes(reports) == ["CALENDAR_NOT_FOUND"]


def test_calendar_uses_base_year_and_reports_bad_rows(tmp_path):
    csv_path = tmp_path / "cal.csv"
    csv_path.write_text("날짜,일정명\n7/7,학원휴무\n이상한값,x\n# 주석,\n", encoding="utf-8")
    events, reports = load_academic_calendar(csv_path, base_year=2027)
    assert list(events) == ["2027-07-07"]
    assert codes(reports) == ["CALENDAR_ROW_SKIPPED"]


def test_calendar_empty_and_header_missing_reported(tmp_path):
    empty = tmp_path / "empty.csv"
    empty.write_text("날짜,일정명\n", encoding="utf-8")
    assert codes(load_academic_calendar(empty)[1]) == ["CALENDAR_EMPTY"]
    bad = tmp_path / "bad.csv"
    bad.write_text("date,name\n7/7,x\n", encoding="utf-8")
    assert codes(load_academic_calendar(bad)[1]) == ["CALENDAR_HEADER_MISSING"]


# ── Medium: --no-pptx 실행에서도 필수값 검사가 수행된다 ──

def test_required_field_reported_without_pptx(tmp_path):
    src = write_sheet(tmp_path / "in.xlsx", make_raw(["7/7"], 강좌명=""))
    res = run_pipeline(src, tmp_path / "out", make_pptx=False)
    rows = list(load_workbook(res["validation_report"]).active.values)
    idx = rows[0].index("issue_code")
    assert "REQUIRED_FIELD_EMPTY" in [r[idx] for r in rows[1:]]


def test_required_field_not_duplicated_with_pptx(tmp_path):
    src = write_sheet(tmp_path / "in.xlsx", make_raw(["7/7"], 강좌명=""))
    res = run_pipeline(src, tmp_path / "out", make_pptx=True, teacher_photo_dir=None)
    rows = list(load_workbook(res["validation_report"]).active.values)
    idx = rows[0].index("issue_code")
    assert [r[idx] for r in rows[1:]].count("REQUIRED_FIELD_EMPTY") == 1


# ── Medium: 수식 문자열은 텍스트로 저장, 계산 캐시 없는 수식은 리포트 ──

def test_formula_like_strings_exported_as_text(tmp_path):
    path = export_validation_report([{"severity": "정보", "raw_value": "=1+1", "message": "+x"}], tmp_path)
    ws = load_workbook(path).active
    values = {c.value: c.data_type for c in ws[2] if c.value}
    assert values["=1+1"] == "s" and values["+x"] == "s"


def test_uncached_formula_cell_reported(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "강좌1"
    ws.append(["강사명", '=CONCAT("홍","길동")'])
    ws.append(["회차", "날짜", "수업 주제", "상세 내용", "비고"])
    path = tmp_path / "f.xlsx"
    wb.save(path)
    _, reports = collect_lectures(path)
    assert codes(reports) == ["FORMULA_NOT_CACHED"]
    assert reports[0]["raw_value"] == "B1"


# ── Medium: 템플릿 복제 시 이미지 관계 유지, 참고 슬라이드 제외, 진도 슬롯 수 자동 인식 ──

def _base_lecture(n_progress=1):
    lec, _ = normalize_lecture(make_raw([f"7/{d}" for d in range(1, n_progress + 1)]), 1, 2026)
    return lec


def test_cloned_slide_keeps_template_image(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    photo = next(PHOTO_DIR.glob("*.jpg"))
    slide.shapes.add_picture(str(photo), Cm(1), Cm(1), width=Cm(1))
    slide.shapes.add_textbox(Cm(1), Cm(3), Cm(10), Cm(1)).text = "{{강좌명}}"
    template = tmp_path / "tpl.pptx"
    prs.save(template)
    path, reports = generate_pptx_from_template([_base_lecture()], template, tmp_path / "out.pptx")
    out = Presentation(path)
    assert len(out.slides) == 1
    assert len(out.slides[0].shapes[0].image.blob) == photo.stat().st_size
    assert "PPTX_GENERATION_FAILED" not in codes(reports)


def test_extra_template_slides_removed_and_reported(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(1)).text = "{{강좌명}}"
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(1)).text = "REFERENCE"
    template = tmp_path / "tpl.pptx"
    prs.save(template)
    path, reports = generate_pptx_from_template([_base_lecture()], template, tmp_path / "out.pptx")
    assert len(Presentation(path).slides) == 1
    assert "TEMPLATE_EXTRA_SLIDES_REMOVED" in codes(reports)


def test_progress_capacity_read_from_template(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for side in ("좌", "우"):
        for idx in range(1, 7):
            slide.shapes.add_textbox(Cm(1), Cm(1 + idx), Cm(5), Cm(0.8)).text = f"{{{{진도_{side}_{idx}_날짜}}}} {{{{진도_{side}_{idx}_내용}}}}"
    assert template_progress_capacity(slide) == 6
    assert progress_split(12, 6) == (6, 6)
    lecture = _base_lecture(12)
    assert validate_progress_overflow(lecture, 6) == []
    overflow = validate_progress_overflow(lecture, 5)
    assert len(overflow) == 1 and "07. 11" in overflow[0]["message"]
    template = tmp_path / "tpl.pptx"
    prs.save(template)
    path, reports = generate_pptx_from_template([lecture], template, tmp_path / "out.pptx")
    assert "PROGRESS_OVERFLOW" not in codes(reports)
    assert "UNMAPPED_PLACEHOLDER" not in codes(reports)
    texts = " ".join(sh.text_frame.text for sh in Presentation(path).slides[0].shapes if sh.has_text_frame)
    assert "07. 12" in texts     # 12번째 진도가 실제로 채워졌다


def test_default_template_capacity_is_five():
    prs = Presentation("templates/강의계획서_마스터템플릿.pptx")
    assert template_progress_capacity(prs.slides[0]) == 5
